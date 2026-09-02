"""任务 3 验证脚本：简历解析 + 简历分析 Agent + 上传路由。

说明：
- 需要 DASHSCOPE_API_KEY 已配置（.env），否则 Agent 调用会失败
- 会真实调用通义千问分析简历
"""
import io
import sys

sys.path.insert(0, "backend")

# 1. 依赖检查
import pypdf
import docx
import openai
print("[1] 依赖检查 OK：pypdf / python-docx / openai")

# 2. 简历解析（内存构造一个 docx）
from docx import Document

doc = Document()
doc.add_heading("张三", 0)
doc.add_paragraph("Python 后端工程师，3 年经验，熟悉 FastAPI/Django、MySQL、Redis、Docker。")
doc.add_paragraph("项目：在线商城系统，负责订单模块与支付回调，日订单量 2 万。")
buf = io.BytesIO()
doc.save(buf)

from app.services.resume_parser import extract_text

with open("_tmp_resume.docx", "wb") as f:
    f.write(buf.getvalue())
text = extract_text("_tmp_resume.docx")
print(f"[2] 简历解析 OK：抽取 {len(text)} 字符，预览：{text[:40]!r}")
assert "张三" in text
import os
os.remove("_tmp_resume.docx")

# 3. 简历分析 Agent（真实调用通义千问）
from app.services.agents.resume_agent import ResumeAgent

try:
    result = ResumeAgent().analyze(text)
    print("[3] 简历分析 Agent OK：", result)
    assert isinstance(result, dict)
    assert "skills" in result
except Exception as exc:
    print(f"[3] 简历分析 Agent 失败（请检查 DASHSCOPE_API_KEY）：{exc}")
    sys.exit(1)

# 4. 上传路由（TestClient，需登录）
from fastapi.testclient import TestClient
from app.main import app

with TestClient(app) as client:
    # 注册拿 token
    client.post("/api/auth/register", json={"username": "task3", "password": "test123456"})
    r = client.post("/api/auth/login", json={"username": "task3", "password": "test123456"})
    token = r.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 上传构造的 docx
    files = {"file": ("resume.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    r = client.post("/api/resume/upload", headers=headers, files=files)
    print(f"[4] 上传接口: {r.status_code}")
    assert r.status_code == 200, r.text
    data = r.json()
    print("    resume_id:", data.get("resume_id"))
    print("    analysis:", data.get("analysis"))
    print("    text_preview:", (data.get("text_preview") or "")[:50])
    assert data.get("resume_id")

    # 5. 重复上传走缓存
    r2 = client.post("/api/resume/upload", headers=headers, files=files)
    assert r2.status_code == 200
    print(f"[5] 重复上传(from_cache={r2.json().get('from_cache')}) OK")

print("== 任务 3 验证全部通过 ==")
